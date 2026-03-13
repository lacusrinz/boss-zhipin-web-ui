#!/usr/bin/env python3
"""
将现有 Excel 文件导入数据库（修正版）
正确处理有"竞品"列的 Excel 文件
"""
import sys
from pathlib import Path

import openpyxl

# 导入数据库模块
from database import BOSSDatabase


def read_excel_file_safe(file_path: str) -> list:
    """
    安全读取 Excel 文件中的岗位数据，正确对齐列

    Args:
        file_path: Excel 文件路径

    Returns:
        list: 岗位数据列表
    """
    print(f"\n正在读取: {Path(file_path).name}")

    try:
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active

        jobs = []
        errors = 0

        # 读取表头
        headers = [str(cell.value).strip() if cell.value else '' for cell in ws[1]]
        print(f"  表头: {headers}")

        # 检查是否有 "竞品" 列
        has_jingpin = '竞品' in headers

        # 确定列索引
        if has_jingpin:
            # 有竞品列的结构
            col_idx = {
                '序号': 0,
                '竞品': 1,
                '岗位名称': 2,
                '薪资': 3,
                '公司名称': 4,
                '公司链接': 5,
                '工作地点': 6,
                '经验要求': 7,
                '学历要求': 8,
                '采集时间': 9,
                '数据来源': 10
            }
        else:
            # 标准结构
            col_idx = {
                '序号': 0,
                '岗位名称': 1,
                '薪资': 2,
                '公司名称': 3,
                '公司链接': 4,
                '工作地点': 5,
                '经验要求': 6,
                '学历要求': 7,
                '采集时间': 8,
                '数据来源': 9
            }

        # 遍历所有数据行
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            try:
                # 跳过空行
                if not row or all(cell is None for cell in row):
                    continue

                # 提取数据
                if has_jingpin:
                    job_name = str(row[col_idx['岗位名称']]).strip() if row[col_idx['岗位名称']] else ''
                    salary = str(row[col_idx['薪资']]).strip() if row[col_idx['薪资']] else ''
                    company = str(row[col_idx['公司名称']]).strip() if row[col_idx['公司名称']] else ''
                    company_url = str(row[col_idx['公司链接']]).strip() if row[col_idx['公司链接']] else ''
                    location = str(row[col_idx['工作地点']]).strip() if row[col_idx['工作地点']] else ''
                    experience = str(row[col_idx['经验要求']]).strip() if row[col_idx['经验要求']] else ''
                    education = str(row[col_idx['学历要求']]).strip() if row[col_idx['学历要求']] else ''
                    collection_time = str(row[col_idx['采集时间']]).strip() if row[col_idx['采集时间']] else ''
                    source = str(row[col_idx['数据来源']]).strip() if row[col_idx['数据来源']] else ''
                else:
                    job_name = str(row[col_idx['岗位名称']]).strip() if row[col_idx['岗位名称']] else ''
                    salary = str(row[col_idx['薪资']]).strip() if row[col_idx['薪资']] else ''
                    company = str(row[col_idx['公司名称']]).strip() if row[col_idx['公司名称']] else ''
                    company_url = str(row[col_idx['公司链接']]).strip() if row[col_idx['公司链接']] else ''
                    location = str(row[col_idx['工作地点']]).strip() if row[col_idx['工作地点']] else ''
                    experience = str(row[col_idx['经验要求']]).strip() if row[col_idx['经验要求']] else ''
                    education = str(row[col_idx['学历要求']]).strip() if row[col_idx['学历要求']] else ''
                    collection_time = str(row[col_idx['采集时间']]).strip() if row[col_idx['采集时间']] else ''
                    source = str(row[col_idx['数据来源']]).strip() if row[col_idx['数据来源']] else ''

                # 验证必要字段
                if not job_name or job_name == 'None':
                    errors += 1
                    continue

                if not company or company == 'None':
                    errors += 1
                    continue

                # 验证薪资不是企业名（企业名不包含K或薪）
                if company and ('K' in company or '薪' in company):
                    errors += 1
                    continue

                job = {
                    'job_name': job_name,
                    'salary': salary,
                    'company': company,
                    'company_url': company_url,
                    'location': location,
                    'experience': experience,
                    'education': education,
                    'source': source if source else Path(file_path).stem.split('_')[0],
                    'collection_time': collection_time
                }

                jobs.append(job)

            except Exception as e:
                errors += 1
                continue

        print(f"  成功: {len(jobs)} 条, 跳过: {errors} 条")
        return jobs

    except Exception as e:
        print(f"  错误: {e}")
        import traceback
        traceback.print_exc()
        return []


def import_all_excels(db: BOSSDatabase, reports_dir: str):
    """导入指定目录下的所有 Excel 文件"""
    reports_path = Path(reports_dir)

    if not reports_path.exists():
        print(f"错误: 目录不存在: {reports_dir}")
        return

    # 查找所有 Excel 文件
    excel_files = sorted(reports_path.glob('*.xlsx'))

    if not excel_files:
        print(f"未找到 Excel 文件: {reports_dir}")
        return

    print(f"\n找到 {len(excel_files)} 个 Excel 文件:")
    for f in excel_files:
        print(f"  - {f.name}")

    print("\n" + "=" * 60)
    print("开始导入数据")
    print("=" * 60)

    total_jobs = []

    for excel_file in excel_files:
        jobs = read_excel_file_safe(str(excel_file))
        if jobs:
            total_jobs.extend(jobs)

    if not total_jobs:
        print("\n没有数据需要导入")
        return

    # 显示预览
    print("\n" + "=" * 60)
    print("数据预览 (前 5 条):")
    print("=" * 60)
    for i, job in enumerate(total_jobs[:5], 1):
        print(f"{i}. 岗位: {job['job_name']}")
        print(f"   企业: {job['company']}")
        print(f"   薪资: {job['salary']}")
        print()

    # 插入数据库
    print("=" * 60)
    print("正在插入数据库...")
    print("=" * 60)

    company_count, job_count = db.insert_jobs_batch(total_jobs)

    print(f"\n✓ 导入完成:")
    print(f"  企业: {company_count} 家")
    print(f"  岗位: {job_count} 个")


def main():
    """主函数"""
    print("=" * 60)
    print("BOSS 直聘 Excel 数据导入工具 (修正版)")
    print("=" * 60)

    # 配置路径
    base_dir = Path(__file__).parent.parent
    reports_dir = base_dir / "reports"

    # 初始化数据库
    db_path = base_dir / "data" / "boss_jobs.db"
    db = BOSSDatabase(str(db_path))

    if not db.connect():
        print("数据库连接失败")
        return 1

    if not db.init_tables():
        print("数据库初始化失败")
        db.close()
        return 1

    # 导入数据
    import_all_excels(db, str(reports_dir))

    # 显示数据库统计
    print("\n" + "=" * 60)
    print("导入后数据库统计")
    print("=" * 60)

    stats = db.get_total_stats()
    print(f"\n企业总数: {stats.get('total_companies', 0)}")
    print(f"岗位总数: {stats.get('total_jobs', 0)}")
    print(f"有 URL 的企业: {stats.get('companies_with_url', 0)}")

    # 来源统计
    source_stats = db.get_source_stats()
    if source_stats:
        print(f"\n按来源统计:")
        for source, count in source_stats:
            print(f"  {source}: {count} 个岗位")

    # 验证企业名
    print("\n" + "=" * 60)
    print("验证企业名")
    print("=" * 60)

    db.cursor.execute("SELECT COUNT(*) FROM companies WHERE name LIKE '%K%' OR name LIKE '%薪%'")
    bad_count = db.cursor.fetchone()[0]

    if bad_count > 0:
        print(f"\n警告: 发现 {bad_count} 个可能包含薪资的企业名")
    else:
        print("\n✓ 企业名验证通过，无薪资格式")

    # 关闭数据库
    db.close()

    print("\n" + "=" * 60)
    print("✓ 导入完成")
    print("=" * 60)

    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("\n\n用户中断程序")
        sys.exit(130)
    except Exception as e:
        print(f"\n程序执行出错: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
