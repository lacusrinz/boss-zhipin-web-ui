#!/usr/bin/env python3
"""
BOSS 直聘岗位信息提取脚本 V6
批量处理 data/collected/ 目录中的 HTML 文件
整合提取、去重、数据库存储、Excel 导出功能
"""

import os
import re
import sys
from datetime import datetime
from pathlib import Path
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from bs4 import BeautifulSoup

# 导入数据库模块
from database import BOSSDatabase


# ==================== 核心提取函数（复用 v5 逻辑）====================

def clean_salary(salary):
    """清理薪资中的特殊字符，将 BOSS 直聘的数字字体图标转换为正常数字"""
    if not salary:
        return ''

    # BOSS 直聘的 Unicode 私有区数字映射 (E031-E03A -> 1-0)
    digit_map = {
        '\ue031': '1',
        '\ue032': '2',
        '\ue033': '3',
        '\ue034': '4',
        '\ue035': '5',
        '\ue036': '6',
        '\ue037': '7',
        '\ue038': '8',
        '\ue039': '9',
        '\ue03a': '0',
    }

    # 替换私有区数字字符
    for char, digit in digit_map.items():
        salary = salary.replace(char, digit)

    # 清理其他私有区字符
    salary = re.sub(r'[\ue000-\uf8ff]', '', salary)

    return salary.strip()


def extract_jobs_from_html(file_path):
    """从 HTML 文件中提取岗位信息（使用 BeautifulSoup）

    Args:
        file_path: HTML 文件路径

    Returns:
        list: 岗位信息列表
    """
    print(f"正在处理: {file_path}")

    # 从文件名提取行业名称
    filename = Path(file_path).stem
    # 文件名格式可能是:
    # - {keyword}_{timestamp}.html (自动收集)
    # - {keyword}.html (手动保存)
    parts = filename.split('_')
    if len(parts) >= 2:
        # 检查最后一部分是否是时间戳
        last_part = parts[-1]
        if last_part.isdigit() and len(last_part) >= 12:
            # 是时间戳格式
            source = '_'.join(parts[:-1])
        else:
            # 不是时间戳，整个文件名就是行业名
            source = filename
    else:
        source = filename  # 使用完整文件名作为来源

    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    soup = BeautifulSoup(content, 'html.parser')
    jobs = []

    # 查找所有岗位卡片
    cards = soup.find_all('li', class_='job-card-box')
    print(f"  找到 {len(cards)} 个岗位卡片")

    for card in cards:
        try:
            # 提取岗位名称
            job_name_tag = card.find('a', class_='job-name')
            job_name = job_name_tag.get_text(strip=True) if job_name_tag else ''

            # 提取薪资
            salary_tag = card.find('span', class_='job-salary')
            salary = clean_salary(salary_tag.get_text(strip=True) if salary_tag else '')

            # 提取公司名称
            company_tag = card.find('span', class_='boss-name')
            company = company_tag.get_text(strip=True) if company_tag else ''

            # 提取公司 URL
            company_url = ''
            boss_info_tag = card.find('a', class_='boss-info')
            if boss_info_tag and boss_info_tag.get('href'):
                href = boss_info_tag.get('href')
                if href and not href.startswith('javascript:'):
                    company_url = f"https://www.zhipin.com{href}"

            # 提取地点
            location_tag = card.find('span', class_='company-location')
            location = location_tag.get_text(strip=True) if location_tag else ''

            # 提取经验和学历
            tags = card.find_all('li')
            experience = ''
            education = ''

            for tag in tags:
                text = tag.get_text(strip=True)
                if not text:
                    continue

                if any(keyword in text for keyword in ['经验', '年', '不限', '应届', '在校']):
                    experience = text
                elif any(keyword in text for keyword in ['学历', '本科', '大专', '硕士', '博士']):
                    education = text

            if job_name:
                job = {
                    'job_name': job_name,
                    'salary': salary,
                    'company': company,
                    'company_url': company_url,
                    'location': location,
                    'experience': experience,
                    'education': education,
                    'source': source
                }
                jobs.append(job)

        except Exception as e:
            print(f"  处理岗位卡片时出错: {e}")
            continue

    print(f"  成功提取 {len(jobs)} 个岗位")
    return jobs


# ==================== 数据处理函数 ====================

def deduplicate_jobs(all_jobs):
    """去重（基于岗位名称和公司名称）"""
    seen = set()
    unique_jobs = []

    for job in all_jobs:
        key = f"{job.get('job_name')}_{job.get('company')}"
        if key not in seen and job.get('job_name'):
            seen.add(key)
            unique_jobs.append(job)

    duplicate_count = len(all_jobs) - len(unique_jobs)
    if duplicate_count > 0:
        print(f"\n去重: 去除 {duplicate_count} 个重复岗位")

    return unique_jobs


# ==================== Excel 导出函数 ====================

def save_to_excel(all_jobs, output_file, collection_time):
    """保存到 Excel 文件"""
    print(f"\n正在保存到 Excel: {output_file}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "岗位信息"

    headers = ['序号', '岗位名称', '薪资', '公司名称', '公司链接', '工作地点',
               '经验要求', '学历要求', '采集时间', '数据来源']

    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for idx, job in enumerate(all_jobs, 1):
        row = idx + 1
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=job.get('job_name', ''))
        ws.cell(row=row, column=3, value=job.get('salary', ''))
        # 兼容两种格式：company 或 company_name
        company = job.get('company_name') or job.get('company', '')
        company_url = job.get('company_url', '')
        ws.cell(row=row, column=4, value=company)
        ws.cell(row=row, column=5, value=company_url)
        ws.cell(row=row, column=6, value=job.get('location', ''))
        ws.cell(row=row, column=7, value=job.get('experience', ''))
        ws.cell(row=row, column=8, value=job.get('education', ''))
        ws.cell(row=row, column=9, value=collection_time)
        ws.cell(row=row, column=10, value=job.get('source', ''))

        for col in range(1, 11):
            ws.cell(row=row, column=col).alignment = Alignment(
                horizontal='center', vertical='center', wrap_text=True
            )

    column_widths = [6, 35, 15, 30, 50, 25, 15, 12, 20, 15]
    for idx, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width

    ws.freeze_panes = 'A2'

    wb.save(output_file)
    print(f"✓ 保存成功，共 {len(all_jobs)} 条记录")


# ==================== 主函数 ====================

def main():
    """主函数：批量处理 collected 目录中的 HTML 文件"""
    # 配置路径
    base_dir = Path(__file__).parent.parent
    collected_dir = base_dir / 'data' / 'collected'
    output_dir = base_dir / 'reports'

    # 创建输出目录
    output_dir.mkdir(exist_ok=True)

    # 采集时间
    collection_time = datetime.now().strftime('%Y-%m-%d %H:%M')

    print("=" * 60)
    print("BOSS 直聘岗位信息批量提取工具 V6")
    print("(支持数据库存储 + Excel 导出)")
    print("=" * 60)

    # 检查 collected 目录是否存在
    if not collected_dir.exists():
        print(f"\n错误: 收集目录不存在: {collected_dir}")
        print("请先手动保存 HTML 文件到 data/collected/ 目录")
        return 1

    # 查找所有 HTML 文件
    html_files = list(collected_dir.glob('*.html'))

    if not html_files:
        print(f"\n错误: 未找到 HTML 文件: {collected_dir}")
        print("请先手动保存 HTML 文件到 data/collected/ 目录")
        return 1

    print(f"\n找到 {len(html_files)} 个 HTML 文件:")
    for f in html_files:
        print(f"  - {f.name}")

    # 提取所有岗位信息
    print("\n" + "=" * 60)
    print("开始提取岗位信息")
    print("=" * 60)

    all_jobs = []
    source_stats = defaultdict(int)

    for html_file in html_files:
        jobs = extract_jobs_from_html(html_file)
        all_jobs.extend(jobs)

        # 统计来源
        if jobs:
            source = jobs[0].get('source', '未知')
            source_stats[source] = len(jobs)

    # 去重
    print("\n" + "=" * 60)
    print("数据去重")
    print("=" * 60)
    unique_jobs = deduplicate_jobs(all_jobs)

    # 初始化数据库
    print("\n" + "=" * 60)
    print("初始化数据库")
    print("=" * 60)

    db_path = base_dir / "data" / "boss_jobs.db"
    db = BOSSDatabase(str(db_path))

    if not db.connect():
        print("警告: 数据库连接失败，仅生成 Excel 报告")
        db = None
    else:
        if not db.init_tables():
            print("警告: 数据库初始化失败，仅生成 Excel 报告")
            db.close()
            db = None
        else:
            print("✓ 数据库初始化成功")

    # 插入数据库
    if db:
        print("\n" + "=" * 60)
        print("正在插入数据库...")
        print("=" * 60)

        # 为每个岗位添加采集时间
        for job in unique_jobs:
            job['collection_time'] = collection_time

        company_count, job_count = db.insert_jobs_batch(unique_jobs)

        print(f"\n✓ 数据库插入完成:")
        print(f"  企业: {company_count} 家")
        print(f"  岗位: {job_count} 个")

        # 显示数据库统计
        print("\n" + "=" * 60)
        print("数据库统计")
        print("=" * 60)

        stats = db.get_total_stats()
        print(f"\n企业总数: {stats.get('total_companies', 0)}")
        print(f"岗位总数: {stats.get('total_jobs', 0)}")
        print(f"有 URL 的企业: {stats.get('companies_with_url', 0)}")

        # 来源统计
        source_stats_db = db.get_source_stats()
        if source_stats_db:
            print(f"\n按来源统计:")
            for source, count in source_stats_db:
                print(f"  {source}: {count} 个岗位")

        # 关闭数据库
        db.close()

    # 输出统计信息
    print("\n" + "=" * 60)
    print("本次提取统计")
    print("=" * 60)

    for source, count in sorted(source_stats.items()):
        print(f"  {source}: {count} 条")

    print(f"\n原始记录: {len(all_jobs)} 条")
    print(f"去重后: {len(unique_jobs)} 条")

    # 保存到 Excel
    timestamp = datetime.now().strftime('%Y%m%d-%H%M%S')
    output_file = output_dir / f'BOSS直聘岗位信息_全行业_{timestamp}.xlsx'

    save_to_excel(unique_jobs, str(output_file), collection_time)

    print(f"\n输出文件: {output_file}")
    print("\n" + "=" * 60)
    print("✓ 提取完成")
    print("=" * 60)

    if db:
        print("\n提示: 使用 'python code/query_database.py' 查询数据库")

    return 0


if __name__ == '__main__':
    sys.exit(main())
